from common.ui import DemoStopped
import io
from datetime import date
import numpy as np
import pandas as pd
import plotly.express as px
import streamlit as st

try:
    from docx import Document
    DOCX_AVAILABLE = True
except Exception:
    Document = None
    DOCX_AVAILABLE = False

st.set_page_config(page_title="Flexible CoA Generator", page_icon="📜", layout="wide")

REQUIRED = {"Part","Characteristic","Value","LSL","Target","USL"}

@st.cache_data
def demo_data():
    rng = np.random.default_rng(42)
    specs = {
        "Length": (9.5,10.0,10.5,"mm"),
        "Weight": (48.0,50.0,52.0,"g"),
        "Thickness": (1.8,2.0,2.2,"mm"),
    }
    prod = [
        ("LOT-1001","BATCH-A","P1","Bottle 500mL","A","Customer A","PO-1000","WO-5001","M1"),
        ("LOT-1002","BATCH-A","P1","Bottle 500mL","A","Customer A","PO-1000","WO-5001","M2"),
        ("LOT-1003","BATCH-B","P1","Bottle 500mL","A","Customer A","PO-1001","WO-5002","M1"),
        ("LOT-2001","BATCH-C","P2","Bottle 1L","B","Customer B","PO-2000","WO-6001","M1"),
        ("LOT-2002","BATCH-C","P2","Bottle 1L","B","Customer B","PO-2000","WO-6001","M2"),
    ]
    rows=[]
    for lot,batch,part,desc,rev,cust,po,wo,machine in prod:
        for char,(lsl,target,usl,unit) in specs.items():
            sigma=(usl-lsl)/8
            shift=(usl-lsl)*0.2 if lot=="LOT-2002" and char=="Weight" else 0
            vals=rng.normal(target+shift,sigma,25)
            for sample,val in enumerate(vals,1):
                if lot=="LOT-2002" and char=="Weight" and sample in (5,18):
                    val=usl+0.15
                rows.append({
                    "Lot":lot,"Batch":batch,"Part":part,"Part Description":desc,
                    "Revision":rev,"Customer":cust,"Purchase Order":po,
                    "Work Order":wo,"Machine":machine,
                    "Manufacturing Date":str(date.today()),
                    "Characteristic":char,"Sample":sample,"Value":round(float(val),4),
                    "LSL":lsl,"Target":target,"USL":usl,"Unit":unit
                })
    return pd.DataFrame(rows)

@st.cache_data
def load_data(f):
    df = pd.read_csv(f) if f.name.lower().endswith(".csv") else pd.read_excel(f)
    missing = REQUIRED.difference(df.columns)
    if missing:
        raise ValueError("Missing required columns: "+", ".join(sorted(missing)))
    df=df.copy()
    for c in ["Value","LSL","Target","USL"]:
        df[c]=pd.to_numeric(df[c],errors="coerce")
    df=df.dropna(subset=["Value","LSL","Target","USL"])
    for c in ["Lot","Batch","Part Description","Revision","Customer","Purchase Order","Work Order","Machine","Manufacturing Date","Unit"]:
        if c not in df.columns:
            df[c]=""
    return df

def cap(group):
    v=group["Value"].dropna()
    mean=v.mean(); std=v.std(ddof=1)
    lsl=float(group["LSL"].iloc[0]); target=float(group["Target"].iloc[0]); usl=float(group["USL"].iloc[0])
    if len(v)<2 or pd.isna(std) or std<=0:
        cp=cpk=np.nan
    else:
        cp=(usl-lsl)/(6*std)
        cpk=min((usl-mean)/(3*std),(mean-lsl)/(3*std))
    return pd.Series({
        "Sample Count":len(v),"Result":mean,"Minimum":v.min(),"Maximum":v.max(),
        "StdDev":std,"Cp":cp,"Cpk":cpk,"LSL":lsl,"Target":target,"USL":usl,
        "Unit":group["Unit"].iloc[0] if "Unit" in group.columns else ""
    })

def coa_table(df):
    t=df.groupby("Characteristic",dropna=False).apply(cap).reset_index()
    t["Status"]=np.where((t["Minimum"]>=t["LSL"])&(t["Maximum"]<=t["USL"]),"PASS","FAIL")
    return t

def uniq(df,col):
    if col not in df.columns: return ""
    vals=df[col].dropna().astype(str).str.strip()
    vals=vals[vals!=""].drop_duplicates().tolist()
    return ", ".join(vals)

def summary(coa,min_cpk):
    fail=coa[coa["Status"]=="FAIL"]
    low=coa[coa["Cpk"].notna()&(coa["Cpk"]<min_cpk)]
    if not fail.empty:
        return "FAIL: out-of-spec measurements found for "+", ".join(fail["Characteristic"].astype(str))+"."
    if not low.empty:
        return f"PASS to specification, but Cpk is below {min_cpk:.2f} for "+", ".join(low["Characteristic"].astype(str))+"."
    return f"PASS: all selected measurements are within specification and available Cpk values meet or exceed {min_cpk:.2f}."


def product_lot_information(df):
    fields = [
        ("Part Number", "Part"),
        ("Part Description", "Part Description"),
        ("Revision", "Revision"),
        ("Customer", "Customer"),
        ("Purchase Order", "Purchase Order"),
        ("Work Order", "Work Order"),
        ("Lot", "Lot"),
        ("Batch", "Batch"),
        ("Machine", "Machine"),
        ("Manufacturing Date", "Manufacturing Date"),
        ("Expiration Date", "Expiration Date"),
    ]

    rows = []
    for label, column in fields:
        if column in df.columns:
            value = uniq(df, column)
            if value != "":
                rows.append({"Field": label, "Value": value})

    return pd.DataFrame(rows)

def make_excel(selected,coa,trace_fields):
    out=io.BytesIO()
    trace=pd.DataFrame([{"Trace Field":f,"Included Values":uniq(selected,f)} for f in trace_fields])
    with pd.ExcelWriter(out,engine="openpyxl") as writer:
        trace.to_excel(writer,sheet_name="Traceability",index=False)
        coa.to_excel(writer,sheet_name="Certificate Results",index=False)
        selected.to_excel(writer,sheet_name="Source Measurements",index=False)
    out.seek(0)
    return out


def make_word(selected, coa, trace_fields, company_name, company_address,
              certificate_number, prepared_by, approved_by,
              compliance_statement, min_cpk):
    if not DOCX_AVAILABLE:
        return None

    doc = Document()
    doc.add_heading("Certificate of Analysis", 0)

    if company_name:
        doc.add_paragraph(company_name)
    if company_address:
        doc.add_paragraph(company_address)

    doc.add_paragraph(f"Certificate Number: {certificate_number}")

    doc.add_heading("Product / Lot Information", level=1)

    info_df = product_lot_information(selected)

    if not info_df.empty:
        info_table = doc.add_table(rows=1, cols=2)
        info_table.style = "Table Grid"
        info_table.rows[0].cells[0].text = "Field"
        info_table.rows[0].cells[1].text = "Value"

        for _, row in info_df.iterrows():
            cells = info_table.add_row().cells
            cells[0].text = str(row["Field"])
            cells[1].text = str(row["Value"])

    doc.add_heading("Traceability / Production Scope", level=1)
    t = doc.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    t.rows[0].cells[0].text = "Trace Field"
    t.rows[0].cells[1].text = "Included Values"

    for field in trace_fields:
        row = t.add_row().cells
        row[0].text = field
        row[1].text = uniq(selected, field)

    doc.add_heading("Test Results", level=1)
    cols = ["Characteristic","LSL","Target","USL","Result","Minimum","Maximum","Unit","Status"]
    rt = doc.add_table(rows=1, cols=len(cols))
    rt.style = "Table Grid"
    for i, col in enumerate(cols):
        rt.rows[0].cells[i].text = col
    for _, row in coa[cols].round(4).iterrows():
        cells = rt.add_row().cells
        for i, col in enumerate(cols):
            cells[i].text = str(row[col])

    doc.add_heading("Capability Summary", level=1)
    cap_cols = ["Characteristic","Sample Count","Result","StdDev","Cp","Cpk"]
    ct = doc.add_table(rows=1, cols=len(cap_cols))
    ct.style = "Table Grid"
    for i, col in enumerate(cap_cols):
        ct.rows[0].cells[i].text = col
    for _, row in coa[cap_cols].round(4).iterrows():
        cells = ct.add_row().cells
        for i, col in enumerate(cap_cols):
            cells[i].text = str(row[col])

    doc.add_heading("Quality Summary", level=1)
    doc.add_paragraph(summary(coa, min_cpk))

    doc.add_heading("Compliance Statement", level=1)
    doc.add_paragraph(compliance_statement)

    doc.add_heading("Approval", level=1)
    doc.add_paragraph(f"Prepared By: {prepared_by}")
    doc.add_paragraph(f"Approved By: {approved_by}")
    doc.add_paragraph(
        "Overall Status: " + ("PASS" if (coa["Status"]=="PASS").all() else "FAIL")
    )

    out = io.BytesIO()
    doc.save(out)
    out.seek(0)
    return out

st.title("📜 Flexible Certificate of Analysis Generator")
from common.ui import render_guide
render_guide('coa')

with st.expander("❓ How This App Works (Click to Expand)", expanded=False):
    st.markdown('### What this app does\nBuilds an example Certificate of Analysis (CoA) from selected quality measurements. It helps demonstrate how results and specifications can be assembled into a consistent report.\n\n### Step 1: Select the report scope\nChoose the trace fields that identify the records you want, then select their values. Trace fields may represent a lot, purchase order, machine, or other identifying context. Check the matching records: a certificate can include more than one lot when its selected scope does. **Clear all trace selections** resets these choices.\n\n### Step 2: Fill in report details\nEnter company information, certificate number, preparer, approver, and compliance statement. These are editable report text, not verified signatures or authorization. The public demo has no data or logo upload.\n\n### Step 3: Review the quality summary\nCheck the selected characteristics, measurements, targets, and lower/upper specification limits. Cpk compares estimated spread and centering with the specification limits. Set the minimum acceptable Cpk for the demonstration and review its effect. This threshold does not redefine specification limits or prove process stability.\n\n### Step 4: Preview and export\nUse the chart-preview option if helpful. Review the complete scope and wording before downloading the Word or Excel CoA. Expand the existing traceability explanation for details about combining selected records.\n\n### Before production use\nThese downloads contain synthetic example results and are not shipment-release evidence. Specifications for a characteristic must be consistent because grouped summaries use the first specification values. A customized implementation needs agreed units, lot traceability, approved wording, and an appropriate review/release process.\n\n### Demo boundaries and customization\nAll analysis uses built-in synthetic data. Explore the controls freely, then use the contact form at the bottom to describe your process and customization needs. Results are examples, not validated production decisions.\n')
st.caption("Create one CoA from any combination of trace fields and multiple selected values.")

with st.sidebar:
    source = 'Demo data'  # Demo-only deployment
    data = demo_data()

excluded={"Value","LSL","Target","USL","Characteristic","Sample","Unit"}
trace_candidates=[c for c in data.columns if c not in excluded]

st.subheader("1. Choose traceability fields")

# Explicitly initialize trace selections as empty for this app session.
if "coa_trace_state_initialized_v3" not in st.session_state:
    st.session_state["trace_fields_v3"] = []
    st.session_state["coa_trace_state_initialized_v3"] = True

if st.button("Clear all trace selections", use_container_width=False, help='Reset the selected trace fields and their filter values so you can choose a new certificate scope.'):
    st.session_state["trace_fields_v3"] = []
    for key in list(st.session_state.keys()):
        if key.startswith("trace_values_v3_"):
            st.session_state[key] = []
    st.rerun()

trace_fields = st.multiselect(
    "The CoA can be filtered by any of these fields",
    trace_candidates,
    key="trace_fields_v3"
, help='Choose identifying fields to filter the certificate records. Then select values for each field below.')

selected = data.copy()

if trace_fields:
    st.subheader("2. Select one or multiple values")

    for field in trace_fields:
        opts = sorted(
            selected[field]
            .dropna()
            .astype(str)
            .str.strip()
            .replace("", np.nan)
            .dropna()
            .unique()
        )

        value_key = "trace_values_v3_" + field

        # Any newly displayed trace-value selector starts empty.
        if value_key not in st.session_state:
            st.session_state[value_key] = []

        # Remove stale values that no longer exist in the available options.
        st.session_state[value_key] = [
            value
            for value in st.session_state[value_key]
            if value in opts
        ]

        vals = st.multiselect(
            field,
            opts,
            key=value_key
        , help='Select the values to include for this trace field. Multiple values can be combined within the certificate scope.')

        if vals:
            selected = selected[
                selected[field].astype(str).isin(vals)
            ]

if selected.empty:
    st.warning("No rows match the selected scope.")
    raise DemoStopped()

st.sidebar.header("Certificate Settings")

logo_file = None  # Public demo has no upload control

company_name = st.sidebar.text_input(
    "Company name",
    value="Your Company Name"
, help='Company name printed on the demonstration certificate. Enter the intended report text.')

company_address = st.sidebar.text_area(
    "Company address",
    value="123 Manufacturing Drive\nCity, State ZIP",
    height=80
, help='Address printed on the demonstration certificate.')

certificate_number = st.sidebar.text_input(
    "Certificate number",
    value="COA-001"
, help='Identifier printed on this certificate. This field does not automatically enforce uniqueness.')

min_cpk = st.sidebar.number_input(
    "Minimum acceptable Cpk",
    min_value=0.50,
    max_value=3.00,
    value=1.33,
    step=0.01
, help='Capability threshold used for evaluation. Cpk measures centering and spread relative to specifications; this is not a stability test.')

show_charts = st.sidebar.checkbox(
    "Include charts in preview",
    value=True
, help='Show charts in the on-screen certificate preview.')

prepared_by = st.sidebar.text_input(
    "Prepared by",
    value="Quality Department"
, help='Name printed as the preparer. This text is not an authenticated signature.')

approved_by = st.sidebar.text_input(
    "Approved by",
    value="Quality Manager"
, help='Name printed as the approver. Entering a name does not record a verified approval.')

compliance_statement = st.sidebar.text_area(
    "Compliance statement",
    value=(
        "This production data has been inspected according to the approved "
        "quality plan and the results shown on this certificate represent "
        "the recorded inspection data for the selected traceability scope."
    ),
    height=120
, help='Wording included on the certificate. Review it against the selected results before issuing any real report.')

coa=coa_table(selected)
status="PASS" if (coa["Status"]=="PASS").all() else "FAIL"

st.subheader("3. Certificate Preview")

if logo_file is not None:
    st.image(logo_file, width=220)

st.markdown("## Certificate of Analysis")
st.markdown(f"**{company_name}**")
if company_address:
    st.write(company_address)
st.write(f"**Certificate Number:** {certificate_number}")

c1,c2,c3,c4=st.columns(4)
c1.metric("Measurements",f"{len(selected):,}")
c2.metric("Characteristics",len(coa))
c3.metric("Trace Fields",len(trace_fields))
c4.metric("Overall Status",status)

st.markdown("### Product / Lot Information")

product_info_df = product_lot_information(selected)

if not product_info_df.empty:
    st.dataframe(
        product_info_df,
        use_container_width=True,
        hide_index=True
    )

st.markdown("### Traceability Scope")
trace_df=pd.DataFrame([{"Trace Field":f,"Included Values":uniq(selected,f)} for f in trace_fields])

if not trace_df.empty:
    st.dataframe(trace_df,use_container_width=True,hide_index=True)
else:
    st.info("No traceability fields selected.")

st.markdown("### Test Results")
cols=["Characteristic","LSL","Target","USL","Result","Minimum","Maximum","Unit","Status"]
st.dataframe(coa[cols].round(4),use_container_width=True,hide_index=True)

st.markdown("### Capability Summary")
st.dataframe(coa[["Characteristic","Sample Count","Result","StdDev","Cp","Cpk"]].round(4),
             use_container_width=True,hide_index=True)

st.markdown("### Quality Summary")
st.write(summary(coa,min_cpk))

if show_charts:
    st.markdown("### Capability Chart")
    fig=px.bar(coa,x="Characteristic",y="Cpk",title="Capability by Characteristic")
    fig.add_hline(y=min_cpk,line_dash="dash",annotation_text=f"Cpk {min_cpk:.2f}")
    st.plotly_chart(fig,use_container_width=True,key="coa_capability_chart")

    st.markdown("### Measurement Distributions")
    for characteristic in coa["Characteristic"]:
        d = selected[selected["Characteristic"] == characteristic].copy()
        if d.empty:
            continue

        hist = px.histogram(
            d,
            x="Value",
            nbins=25,
            title=f"{characteristic} Distribution"
        )
        hist.add_vline(
            x=float(d["LSL"].iloc[0]),
            line_dash="dot",
            annotation_text="LSL"
        )
        hist.add_vline(
            x=float(d["Target"].iloc[0]),
            line_dash="dash",
            annotation_text="Target"
        )
        hist.add_vline(
            x=float(d["USL"].iloc[0]),
            line_dash="dot",
            annotation_text="USL"
        )
        st.plotly_chart(
            hist,
            use_container_width=True,
            key=f"distribution_{characteristic}"
        )

st.markdown("### Compliance Statement")
st.write(compliance_statement)

st.markdown("### Approval")
a1, a2 = st.columns(2)
a1.write(f"**Prepared By:** {prepared_by}")
a2.write(f"**Approved By:** {approved_by}")

st.markdown("### Downloads")
d1, d2 = st.columns(2)

with d1:
    if DOCX_AVAILABLE:
        st.download_button(
            "Download Word CoA",
            data=make_word(
                selected,
                coa,
                trace_fields,
                company_name,
                company_address,
                certificate_number,
                prepared_by,
                approved_by,
                compliance_statement,
                min_cpk
            ),
            file_name=f"{certificate_number}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True
        , help='Download the demonstration certificate as a Word document for review.')
    else:
        st.info("Word export unavailable. Install python-docx to enable it.")

with d2:
    st.download_button(
        "Download Excel CoA",
        data=make_excel(selected,coa,trace_fields),
        file_name=f"{certificate_number}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True
    , help='Download the certificate results in Excel format using the current demonstration scope.')

with st.expander("How traceability works"):
    st.markdown("""
The CoA is not tied to Lot.

You can use a trace field from the demonstration data, such as:
- Lot
- Batch
- Purchase Order
- Work Order
- Customer
- Machine
- Part
- Revision
- Manufacturing Date
- Any other identifying column

Each selected field can contain multiple values.

Example:
**Purchase Order PO-1000** may include **LOT-1001 + LOT-1002**, **WO-5001**, and **M1 + M2**.
The CoA combines all matching measurements and calculates the result and capability from that selected scope.
""")

with st.expander("Data fields used in a customized version"):
    st.code("Part, Characteristic, Value, LSL, Target, USL")
    st.write("All traceability fields are optional and become selectable automatically when present.")
